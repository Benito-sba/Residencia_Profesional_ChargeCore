from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from bson import ObjectId
from bson.errors import InvalidId
from datetime import datetime, timedelta
from typing import Optional
import io

from app.database import get_collection
from app.auth import require_read_access
from app.validation import parse_object_id

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])


def _fmt_date(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


async def _resolve_user_name(user_id):
    if not user_id:
        return ""
    try:
        object_id = ObjectId(user_id)
    except (InvalidId, TypeError):
        return ""
    user = await get_collection("usuarios").find_one({"_id": object_id})
    return user.get("nombre", "") if user else ""


async def _resolve_vehicle_name(vehicle_id):
    if not vehicle_id:
        return ""
    try:
        object_id = ObjectId(vehicle_id)
    except (InvalidId, TypeError):
        return ""
    vehicle = await get_collection("vehiculos").find_one({"_id": object_id})
    return vehicle.get("modelo", "") if vehicle else ""


async def _resolve_nfc(uid):
    if not uid:
        return ""
    try:
        object_id = ObjectId(uid)
    except (InvalidId, TypeError):
        return str(uid)
    card = await get_collection("nfc_cards").find_one({"_id": object_id})
    return card.get("uid", "") if card else ""


def _parse_report_date(value: str, field_name: str) -> datetime:
    try:
        return datetime.fromisoformat(value)
    except (TypeError, ValueError):
        raise HTTPException(
            status_code=422,
            detail=f"El valor de {field_name} no es una fecha válida",
        )


def _parse_report_end_date(value: str, field_name: str) -> tuple[datetime, str]:
    parsed = _parse_report_date(value, field_name)
    if len(value) == 10:
        return parsed + timedelta(days=1), "$lt"
    return parsed, "$lte"


@router.get("/dashboard")
async def dashboard_report(current_user=Depends(require_read_access)):
    sesiones_col = get_collection("sesiones")
    usuarios_col = get_collection("usuarios")
    vehiculos_col = get_collection("vehiculos")
    cargadores_col = get_collection("cargadores")
    recompensas_col = get_collection("recompensas")

    total_usuarios = await usuarios_col.count_documents({})
    total_vehiculos = await vehiculos_col.count_documents({})
    total_cargadores = await cargadores_col.count_documents({"activo": True})
    cargadores_conectados = await cargadores_col.count_documents({"activo": True, "estado_general": {"$in": ["disponible", "ocupado"]}})
    cargadores_ocupados = await cargadores_col.count_documents({"activo": True, "estado_general": "ocupado"})
    sesiones_realizadas = await sesiones_col.count_documents({"estado": "finalizada"})

    energia_pipeline = [{"$match": {"estado": "finalizada"}}, {"$group": {"_id": None, "total": {"$sum": "$energia_consumida_kwh"}}}]
    energia_result = await sesiones_col.aggregate(energia_pipeline).to_list(1)
    kwh_consumidos = round(energia_result[0]["total"], 2) if energia_result else 0

    recompensas_otorgadas = await recompensas_col.count_documents({})

    consumo_por_mes_pipeline = [
        {"$match": {"estado": "finalizada", "fecha_inicio": {"$gte": datetime.utcnow() - timedelta(days=365)}}},
        {"$group": {"_id": {"year": {"$year": "$fecha_inicio"}, "month": {"$month": "$fecha_inicio"}}, "kwh": {"$sum": "$energia_consumida_kwh"}}},
        {"$sort": {"_id.year": 1, "_id.month": 1}},
    ]
    consumo_mes = []
    async for item in sesiones_col.aggregate(consumo_por_mes_pipeline):
        mes = f"{item['_id']['year']}-{item['_id']['month']:02d}"
        consumo_mes.append({"mes": mes, "kwh": round(item["kwh"], 2)})

    sesiones_por_dia_pipeline = [
        {"$match": {"estado": "finalizada", "fecha_inicio": {"$gte": datetime.utcnow() - timedelta(days=30)}}},
        {"$group": {"_id": {"year": {"$year": "$fecha_inicio"}, "month": {"$month": "$fecha_inicio"}, "day": {"$dayOfMonth": "$fecha_inicio"}}, "total": {"$sum": 1}}},
        {"$sort": {"_id.year": 1, "_id.month": 1, "_id.day": 1}},
    ]
    sesiones_dia = []
    async for item in sesiones_col.aggregate(sesiones_por_dia_pipeline):
        fecha = f"{item['_id']['year']}-{item['_id']['month']:02d}-{item['_id']['day']:02d}"
        sesiones_dia.append({"fecha": fecha, "sesiones": item["total"]})

    uso_cargadores = []
    async for cargador in cargadores_col.find({"activo": True}).sort("charge_point_id", 1):
        sesiones = await sesiones_col.count_documents({"charge_point_id": cargador.get("charge_point_id"), "estado": "finalizada"})
        uso_cargadores.append({"cargador": cargador.get("charge_point_id"), "sesiones": sesiones})

    top_usuarios_pipeline = [
        {"$match": {"estado": "finalizada"}},
        {"$group": {"_id": "$usuario_id", "total_kwh": {"$sum": "$energia_consumida_kwh"}, "sesiones": {"$sum": 1}}},
        {"$sort": {"total_kwh": -1}},
        {"$limit": 5},
    ]
    top_usuarios = []
    async for item in sesiones_col.aggregate(top_usuarios_pipeline):
        nombre = await _resolve_user_name(str(item["_id"])) if item.get("_id") else "Sin usuario"
        top_usuarios.append({"usuario": nombre, "kwh": round(item["total_kwh"], 2), "sesiones": item["sesiones"]})

    recompensas_mes_pipeline = [
        {"$match": {"fecha_otorgamiento": {"$gte": datetime.utcnow() - timedelta(days=365)}}},
        {"$group": {"_id": {"year": {"$year": "$fecha_otorgamiento"}, "month": {"$month": "$fecha_otorgamiento"}}, "total": {"$sum": 1}}},
        {"$sort": {"_id.year": 1, "_id.month": 1}},
    ]
    recompensas_mes = []
    async for item in recompensas_col.aggregate(recompensas_mes_pipeline):
        mes = f"{item['_id']['year']}-{item['_id']['month']:02d}"
        recompensas_mes.append({"mes": mes, "recompensas": item["total"]})

    return {
        "kpis": {
            "usuarios_registrados": total_usuarios,
            "vehiculos_registrados": total_vehiculos,
            "cargadores_conectados": cargadores_conectados,
            "cargadores_ocupados": cargadores_ocupados,
            "sesiones_realizadas": sesiones_realizadas,
            "kwh_consumidos": round(kwh_consumidos, 2),
            "recompensas_otorgadas": recompensas_otorgadas,
        },
        "graficas": {
            "consumo_por_mes": consumo_mes,
            "sesiones_por_dia": sesiones_dia,
            "uso_de_cargadores": uso_cargadores,
            "top_usuarios": top_usuarios,
            "recompensas_generadas": recompensas_mes,
        },
    }


@router.get("/sesiones")
async def report_sessions(
    usuario_id: Optional[str] = None,
    charge_point_id: Optional[str] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    limit: int = 100,
    current_user=Depends(require_read_access),
):
    query = {"estado": "finalizada"}
    if usuario_id:
        query["usuario_id"] = parse_object_id(usuario_id, "usuario")
    if charge_point_id:
        query["charge_point_id"] = charge_point_id
    if fecha_inicio:
        query["fecha_inicio"] = {
            "$gte": _parse_report_date(fecha_inicio, "fecha_inicio")
        }
    if fecha_fin:
        end_date, operator = _parse_report_end_date(fecha_fin, "fecha_fin")
        query.setdefault("fecha_inicio", {})[operator] = end_date

    cursor = get_collection("sesiones").find(query).sort("fecha_inicio", -1).limit(limit)
    rows = []
    async for item in cursor:
        item["_id"] = str(item.get("_id"))
        item["usuario"] = await _resolve_user_name(str(item.get("usuario_id"))) if item.get("usuario_id") else ""
        item["vehiculo"] = await _resolve_vehicle_name(str(item.get("vehiculo_id"))) if item.get("vehiculo_id") else ""
        item["nfc"] = await _resolve_nfc(str(item.get("nfc_card_id"))) if item.get("nfc_card_id") else ""
        item["nfc_card_id"] = str(item.get("nfc_card_id")) if item.get("nfc_card_id") else None
        item["fecha_inicio"] = _fmt_date(item.get("fecha_inicio"))
        item["fecha_fin"] = _fmt_date(item.get("fecha_fin"))
        item["duracion_minutos"] = item.get("duracion_minutos", 0)
        item["energia_consumida_kwh"] = round(item.get("energia_consumida_kwh", 0), 2)
        rows.append(item)
    return rows


@router.get("/consumo")
async def report_consumo(
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    limit: int = 10,
    current_user=Depends(require_read_access),
):
    query = {"estado": "finalizada"}
    if fecha_inicio or fecha_fin:
        range_filter = {}
        if fecha_inicio:
            range_filter["$gte"] = _parse_report_date(fecha_inicio, "fecha_inicio")
        if fecha_fin:
            end_date, operator = _parse_report_end_date(fecha_fin, "fecha_fin")
            range_filter[operator] = end_date
        query["fecha_inicio"] = range_filter

    pipeline = [
        {"$match": query},
        {"$group": {"_id": "$usuario_id", "total_cargas": {"$sum": 1}, "total_kwh": {"$sum": "$energia_consumida_kwh"}, "promedio_kwh": {"$avg": "$energia_consumida_kwh"}}},
        {"$sort": {"total_kwh": -1}},
        {"$limit": limit},
    ]
    rows = []
    async for item in get_collection("sesiones").aggregate(pipeline):
        rows.append({
            "usuario": await _resolve_user_name(str(item["_id"])) if item.get("_id") else "Sin usuario",
            "total_cargas": item["total_cargas"],
            "total_kwh": round(item["total_kwh"], 2),
            "promedio_kwh": round(item["promedio_kwh"], 2),
        })
    return rows


@router.get("/usuarios")
async def report_usuarios(current_user=Depends(require_read_access)):
    rows = []
    async for user in get_collection("usuarios").find({}).sort("created_at", -1):
        vehiculos = await get_collection("vehiculos").count_documents({"usuario_id": user.get("_id")})
        cargas = user.get("total_cargas", 0)
        rows.append({
            "nombre": user.get("nombre", ""),
            "correo": user.get("email", ""),
            "nfc_asociado": user.get("nfc_uid", "N/A"),
            "vehiculos_registrados": vehiculos,
            "numero_cargas": cargas,
            "estado": "Activo" if user.get("activo", True) else "Inactivo",
        })
    return rows


@router.get("/recompensas")
async def report_rewards(current_user=Depends(require_read_access)):
    rows = []
    async for reward in get_collection("recompensas").find({}).sort("fecha_otorgamiento", -1):
        rows.append({
            "usuario": await _resolve_user_name(str(reward.get("usuario_id"))) if reward.get("usuario_id") else "",
            "tipo_de_recompensa": reward.get("tipo", ""),
            "fecha_generacion": _fmt_date(reward.get("fecha_otorgamiento")),
            "estado": "Usada" if reward.get("usado") else "Pendiente",
        })
    return rows


@router.get("/cargadores")
async def report_chargers(current_user=Depends(require_read_access)):
    rows = []
    async for charger in get_collection("cargadores").find({"activo": True}).sort("charge_point_id", 1):
        sesiones = await get_collection("sesiones").count_documents({"charge_point_id": charger.get("charge_point_id"), "estado": "finalizada"})
        energia = await get_collection("sesiones").aggregate([
            {"$match": {"charge_point_id": charger.get("charge_point_id"), "estado": "finalizada"}},
            {"$group": {"_id": None, "total": {"$sum": "$energia_consumida_kwh"}}},
        ]).to_list(1)
        rows.append({
            "charger_id": charger.get("charge_point_id", ""),
            "estado": charger.get("estado_general", ""),
            "ultima_conexion": _fmt_date(charger.get("ultimo_heartbeat")),
            "total_sesiones": sesiones,
            "energia_suministrada_kwh": round(energia[0]["total"], 2) if energia else 0,
        })
    return rows


async def _export_report_excel(rows: list, columns: list, filename: str):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise RuntimeError("openpyxl no está instalado")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte"
    for col, header in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, key in enumerate(columns, start=1):
            value = row.get(key, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    for column in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 30)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}"})


async def _export_report_pdf(rows: list, columns: list, filename: str):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    except ImportError:
        raise RuntimeError("reportlab no está instalado")

    table_data = [columns]
    for row in rows:
        table_data.append([row.get(col, "") for col in columns])

    pdf_buffer = io.BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter)
    table = Table(table_data)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 1, colors.grey),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    doc.build([table])
    pdf_buffer.seek(0)
    return StreamingResponse(pdf_buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/sesiones/export")
async def export_sessions_report(
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    usuario_id: Optional[str] = None,
    charge_point_id: Optional[str] = None,
    fmt: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    current_user=Depends(require_read_access),
):
    rows = await report_sessions(usuario_id, charge_point_id, fecha_inicio, fecha_fin, limit=500, current_user=current_user)
    columns = ["_id", "charge_point_id", "usuario", "vehiculo", "nfc", "fecha_inicio", "fecha_fin", "duracion_minutos", "energia_consumida_kwh"]
    filename = f"sesiones_{datetime.now().strftime('%Y%m%d')}.{fmt}"
    try:
        if fmt == "pdf":
            return await _export_report_pdf(rows, columns, filename)
        return await _export_report_excel(rows, columns, filename)
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.get("/usuarios/export")
async def export_users_report(
    fmt: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    current_user=Depends(require_read_access),
):
    rows = await report_usuarios(current_user=current_user)
    columns = ["nombre", "correo", "nfc_asociado", "vehiculos_registrados", "numero_cargas", "estado"]
    filename = f"usuarios_{datetime.now().strftime('%Y%m%d')}.{fmt}"
    try:
        if fmt == "pdf":
            return await _export_report_pdf(rows, columns, filename)
        return await _export_report_excel(rows, columns, filename)
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.get("/recompensas/export")
async def export_rewards_report(
    fmt: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    current_user=Depends(require_read_access),
):
    rows = await report_rewards(current_user=current_user)
    columns = ["usuario", "tipo_de_recompensa", "fecha_generacion", "estado"]
    filename = f"recompensas_{datetime.now().strftime('%Y%m%d')}.{fmt}"
    try:
        if fmt == "pdf":
            return await _export_report_pdf(rows, columns, filename)
        return await _export_report_excel(rows, columns, filename)
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.get("/cargadores/export")
async def export_chargers_report(
    fmt: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    current_user=Depends(require_read_access),
):
    rows = await report_chargers(current_user=current_user)
    columns = ["charger_id", "estado", "ultima_conexion", "total_sesiones", "energia_suministrada_kwh"]
    filename = f"cargadores_{datetime.now().strftime('%Y%m%d')}.{fmt}"
    try:
        if fmt == "pdf":
            return await _export_report_pdf(rows, columns, filename)
        return await _export_report_excel(rows, columns, filename)
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc


@router.get("/consumo/export")
async def export_consumo_report(
    fmt: str = Query("xlsx", pattern="^(xlsx|pdf)$"),
    current_user=Depends(require_read_access),
):
    rows = await report_consumo(current_user=current_user)
    columns = ["usuario", "total_cargas", "total_kwh", "promedio_kwh"]
    filename = f"consumo_{datetime.now().strftime('%Y%m%d')}.{fmt}"
    try:
        if fmt == "pdf":
            return await _export_report_pdf(rows, columns, filename)
        return await _export_report_excel(rows, columns, filename)
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail=str(exc)) from exc